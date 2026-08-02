from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from datetime import date

import openpyxl
import pytest

from app.domain import AuditKind, EntryStatus, canonical_json
from app.exporter import (
    ExportPreflightError,
    build_generated_division_roster_workbook,
    compare_workbook_cells,
    prepare_generated_division_roster_export,
    save_ngo_division_workbook,
)
from app.importer import parse_division_workbook
from app.scheduler import run_scheduler
from app.services.weekly_demo import WeeklyRosterDemoBuilder
from fixtures.paths import DIVISION_WORKBOOK_PATH
from fixtures.paths import ESCORT_WORKBOOK_PATH, HC_TIMETABLE_WORKBOOK_PATH


def _weekly_demo_result():
    build = WeeklyRosterDemoBuilder().build(
        hc_workbook_path=HC_TIMETABLE_WORKBOOK_PATH,
        escort_workbook_path=ESCORT_WORKBOOK_PATH,
        week_start=date(2026, 1, 5),
        changes_json="[]",
    )
    return build, run_scheduler(build.snapshot)


@pytest.fixture(scope="module")
def prepared_export_plan():
    build, result = _weekly_demo_result()
    plan = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=result.version,
        generated=result.generated,
    )
    return build, result, plan


def test_ngo_format_exporter_preserves_business_sheets_on_no_edit_roundtrip(tmp_path):
    out = save_ngo_division_workbook(
        template_path=DIVISION_WORKBOOK_PATH,
        output_dir=tmp_path,
    )
    wb = openpyxl.load_workbook(out)

    assert {"RC_變更摘要", "RC_審核", "RC_未分配", "RC_meta"} <= set(wb.sheetnames)
    assert compare_workbook_cells(DIVISION_WORKBOOK_PATH, out) == []


def test_changed_cell_marker_uses_border_and_comment_without_replacing_fill(tmp_path):
    original = openpyxl.load_workbook(DIVISION_WORKBOOK_PATH)
    original_fill = original["恆常服務"]["D12"].fill.fgColor.rgb
    out = save_ngo_division_workbook(
        template_path=DIVISION_WORKBOOK_PATH,
        output_dir=tmp_path,
        changed_cells={"恆常服務!D12": "test change"},
    )
    changed = openpyxl.load_workbook(out)["恆常服務"]["D12"]

    assert changed.fill.fgColor.rgb == original_fill
    assert changed.border.left.style == "medium"
    assert changed.comment is not None
    assert "RosterCopiilot" in changed.comment.text


def test_generated_export_preflight_revalidates_and_blocks_invalid_grid_write():
    build, result = _weekly_demo_result()
    bad_version = result.version.model_copy(deep=True)
    target = next(
        entry for entry in bad_version.entries
        if entry.status == EntryStatus.SCHEDULED and entry.worker_id
    )
    target.worker_id = "unknown-worker"
    target.worker_name = "不存在同工"

    report = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=bad_version,
        generated=result.generated,
    ).report

    assert report.review_export_allowed is False
    assert report.publication_state == "blocked"
    assert report.validator_violations
    assert any(item.is_export_failure for item in report.unassigned_items)
    with pytest.raises(ExportPreflightError):
        build_generated_division_roster_workbook(
            template_path=DIVISION_WORKBOOK_PATH,
            division_layout=build.division,
            dataset=result.dataset,
            version=bad_version,
            generated=result.generated,
        )


def test_export_mapper_failure_is_reconciled_into_review_and_unassigned_rows():
    build, result = _weekly_demo_result()
    bad_version = result.version.model_copy(deep=True)
    target = next(
        entry for entry in bad_version.entries
        if entry.status == EntryStatus.SCHEDULED and entry.worker_id
    )
    bad_version.entries = [target]
    bad_version.audit_items = []
    bad_version.unassigned = []
    bad_layout = replace(
        build.division,
        workers=tuple(
            worker for worker in build.division.workers
            if worker.display_name != target.worker_name
        ),
    )

    plan = prepare_generated_division_roster_export(
        division_layout=bad_layout,
        dataset=result.dataset,
        version=bad_version,
        generated=result.generated,
    )

    assert plan.report.review_export_allowed is False
    failures = [
        failure for failure in plan.report.export_failures
        if failure.code == "worker_unmapped" and failure.demand_id == target.demand_id
    ]
    assert len(failures) == 1
    failure_items = [
        item for item in plan.report.unassigned_items
        if item.is_export_failure and item.entry_id == target.id
    ]
    assert len(failure_items) == 1
    failure_item = failure_items[0]
    assert failure_item.demand_id == target.demand_id
    assert failure_item.disposition == "scheduled"
    assert failure_item.audit_id
    linked_audits = [
        audit for audit in plan.review_version.audit_items
        if target.demand_id in audit.demand_ids
        and audit.kind == AuditKind.UNASSIGNED_TASK
        and audit.blocking
    ]
    assert len(linked_audits) == 1
    audit = linked_audits[0]
    assert audit.id == failure_item.audit_id == failures[0].audit_id
    assert audit.entry_ids == [target.id]
    assert audit.id in plan.review_version.entry_by_id(target.id).audit_ids
    disposition = next(
        item for item in plan.report.reconciliation.dispositions
        if item.demand_id == target.demand_id
    )
    assert disposition.disposition == "scheduled"
    assert disposition.entry_id == target.id


def test_export_preflight_detects_full_period_cell_collision(monkeypatch):
    build, result = _weekly_demo_result()
    version = result.version.model_copy(deep=True)
    scheduled = [
        entry for entry in version.entries
        if entry.status == EntryStatus.SCHEDULED and entry.worker_id
    ]
    first = scheduled[0].model_copy(update={
        "id": "collision-1", "revision": 1, "session_index": 1,
    })
    second = first.model_copy(update={
        "id": "collision-2", "revision": 2, "session_index": 2,
    })
    full = first.model_copy(update={
        "id": "collision-full", "revision": 3, "session_index": None,
    })
    version.entries = [first, second, full]
    version.unassigned = []
    version.audit_items = []
    monkeypatch.setattr(
        "app.exporter.division_writer.validate_entries",
        lambda *args, **kwargs: [],
    )

    report = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=version,
        generated=result.generated,
    ).report

    assert report.review_export_allowed is False
    assert any(failure.code == "cell_collision" for failure in report.export_failures)


def test_generated_export_preserves_existing_comment_border_and_fill(tmp_path):
    build, result = _weekly_demo_result()
    layout = parse_division_workbook(DIVISION_WORKBOOK_PATH)
    target = next(
        entry for entry in result.version.entries
        if entry.worker_name == "家偉"
        and entry.weekday == 3
        and entry.period.value == "PM"
        and entry.session_index == 1
        and entry.status == EntryStatus.NEEDS_REVIEW
    )
    assert target.demand_id

    original = openpyxl.load_workbook(DIVISION_WORKBOOK_PATH)["恆常服務"]["F48"]
    wb = build_generated_division_roster_workbook(
        template_path=DIVISION_WORKBOOK_PATH,
        division_layout=layout,
        dataset=result.dataset,
        version=result.version,
        generated=result.generated,
    )
    changed = wb["恆常服務"]["F48"]

    assert changed.value
    assert changed.fill.fgColor.rgb == original.fill.fgColor.rgb
    assert changed.border.top.style == original.border.top.style
    assert changed.border.left.style == "medium"
    assert changed.comment is not None
    assert changed.comment.text.startswith("RC:待審\n原因:")
    assert "6/3 開始改時間" in changed.comment.text
    assert f"需求ID: {target.demand_id}" in changed.comment.text
    assert f"項目ID: {target.id}" in changed.comment.text
    assert "來源證據ID:" in changed.comment.text


def test_unassigned_dispositions_have_one_terminal_audit_and_one_rc_row():
    build, result = _weekly_demo_result()
    plan = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=result.version,
        generated=result.generated,
    )

    audits = {audit.id: audit for audit in plan.review_version.audit_items}
    rc_rows_by_demand = {
        item.demand_id: item
        for item in plan.report.unassigned_items
        if not item.is_export_failure
    }
    unassigned = [
        item for item in plan.report.reconciliation.dispositions
        if item.disposition == "unassigned"
    ]
    assert unassigned
    assert len(rc_rows_by_demand) == len(unassigned)
    for disposition in unassigned:
        terminal = [
            audit_id for audit_id in disposition.audit_ids
            if audits[audit_id].kind
            in {AuditKind.UNASSIGNED_TASK, AuditKind.DUTY_UNDER_COVERAGE}
            and audits[audit_id].blocking
        ]
        assert len(terminal) == 1
        row = rc_rows_by_demand[disposition.demand_id]
        assert row.audit_id == terminal[0]
        assert row.entry_id == disposition.entry_id
        entry = plan.review_version.entry_by_id(disposition.entry_id)
        assert entry is not None
        expected_kind = (
            AuditKind.DUTY_UNDER_COVERAGE
            if entry.service_code.value in {"AMC", "MRC", "GC"}
            else AuditKind.UNASSIGNED_TASK
        )
        assert audits[terminal[0]].kind == expected_kind


def test_export_manifest_has_exact_cells_and_resolvable_canonical_links():
    build, result = _weekly_demo_result()
    plan = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=result.version,
        generated=result.generated,
    )

    entries = {entry.id: entry for entry in plan.review_version.entries}
    audits = {audit.id for audit in plan.review_version.audit_items}
    evidence = {item.id for item in result.generated.source_evidence}
    gaps = {item.id for item in result.generated.data_gaps}
    dispositions = {
        item.demand_id: item for item in plan.report.reconciliation.dispositions
    }
    assert len(plan.report.placements) == plan.report.reconciliation.placement_count
    assert plan.report.changed_cell_count == len(plan.changed_cells)
    for placement in plan.report.placements:
        entry = entries[placement.entry_id]
        disposition = dispositions[placement.demand_id]
        assert placement.version_id == plan.review_version.id
        assert placement.demand_id == entry.demand_id
        assert placement.disposition == disposition.disposition
        assert disposition.entry_id == entry.id
        assert placement.assignment_cell.startswith(f"{build.division.sheet_name}!")
        assert placement.assignment_cell in plan.changed_cells
        if placement.detail_cell:
            assert placement.detail_cell.startswith(f"{build.division.sheet_name}!")
            assert placement.detail_cell in plan.changed_cells
        assert set(placement.audit_ids) <= audits
        assert placement.data_gap_ids == sorted(entry.data_gap_ids)
        assert set(placement.data_gap_ids) <= gaps
        assert placement.source_evidence_ids == sorted(
            item.id for item in entry.source_evidence
        )
        assert set(placement.source_evidence_ids) <= evidence


def test_review_comment_uses_exact_provenance_order_and_canonical_ids():
    build, result = _weekly_demo_result()
    plan = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=result.version,
        generated=result.generated,
    )
    review_placement = next(
        item for item in plan.report.placements
        if item.disposition == "needs_review"
    )
    workbook = build_generated_division_roster_workbook(
        template_path=DIVISION_WORKBOOK_PATH,
        division_layout=build.division,
        dataset=result.dataset,
        version=plan.review_version,
        prepared_plan=plan,
    )
    sheet_name, coordinate = review_placement.assignment_cell.split("!", 1)
    comment = workbook[sheet_name][coordinate].comment
    assert comment is not None
    lines = comment.text.splitlines()
    assert lines[0] == "RC:待審"
    assert lines[1].startswith("原因: ")
    assert lines[2] == f"審核ID: {', '.join(review_placement.audit_ids)}"
    assert lines[3] == f"需求ID: {review_placement.demand_id}"
    assert lines[4] == f"項目ID: {review_placement.entry_id}"
    assert lines[5] == (
        "來源證據ID: " + ", ".join(review_placement.source_evidence_ids)
    )


def test_all_rc_sheets_embed_the_same_canonical_reconciliation():
    build, result = _weekly_demo_result()
    plan = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=result.version,
        generated=result.generated,
    )
    workbook = build_generated_division_roster_workbook(
        template_path=DIVISION_WORKBOOK_PATH,
        division_layout=build.division,
        dataset=result.dataset,
        version=plan.review_version,
        prepared_plan=plan,
    )
    expected = canonical_json(plan.report.reconciliation.model_dump(mode="json"))
    for sheet_name in ("RC_變更摘要", "RC_審核", "RC_未分配", "RC_meta"):
        assert _embedded_reconciliation(workbook[sheet_name]) == expected

    expected_counts = {
        key: value for key, value in _sheet_key_values(workbook["RC_meta"]).items()
        if key.startswith("reconciliation.")
    }
    summary_counts = {
        key: value for key, value in _sheet_key_values(workbook["RC_變更摘要"]).items()
        if key.startswith("reconciliation.")
    }
    assert summary_counts == expected_counts


def test_missing_demand_link_fails_closed_without_exporter_rewriting_it():
    build, result = _weekly_demo_result()
    version = result.version.model_copy(deep=True)
    target = next(
        entry for entry in version.entries
        if entry.status == EntryStatus.SCHEDULED and not entry.audit_ids
    )
    target.demand_id = "dem_00000000000000000000"

    plan = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=version,
        generated=result.generated,
    )

    errors = [
        error for error in plan.report.reconciliation.errors
        if error.code == "missing_demand_link"
        and target.demand_id in error.demand_ids
    ]
    assert errors
    assert plan.report.review_export_allowed is False
    assert not plan.report.export_failures
    assert all(
        target.demand_id not in audit.demand_ids
        or all(reason.code.value != "export_placement_failure" for reason in audit.reasons)
        for audit in plan.review_version.audit_items
    )


def test_needs_review_without_audit_fails_closed_without_ad_hoc_audit():
    build, result = _weekly_demo_result()
    version = result.version.model_copy(deep=True)
    audit_entry_ids = {
        entry_id for audit in version.audit_items for entry_id in audit.entry_ids
    }
    target = next(
        entry for entry in version.entries
        if entry.status == EntryStatus.SCHEDULED
        and not entry.audit_ids
        and entry.id not in audit_entry_ids
        and not entry.data_gap_ids
        and all(item.confidence not in {"low", "seed"} for item in entry.source_evidence)
    )
    target.status = EntryStatus.NEEDS_REVIEW
    target.review_reasons = []
    target.audit_ids = []

    plan = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=version,
        generated=result.generated,
    )

    errors = [
        error for error in plan.report.reconciliation.errors
        if error.code == "missing_audit_link" and target.demand_id in error.demand_ids
    ]
    assert errors
    assert plan.report.review_export_allowed is False
    assert not plan.report.export_failures
    review_entry = next(
        entry for entry in plan.review_version.entries
        if entry.demand_id == target.demand_id
    )
    assert review_entry.audit_ids == []
    assert not [
        audit for audit in plan.review_version.audit_items
        if target.demand_id in audit.demand_ids
    ]


def test_preflight_preserves_caller_version_template_and_legal_blocked_draft():
    build, result = _weekly_demo_result()
    version_before = canonical_json(result.version.model_dump(mode="json"))
    template_before = DIVISION_WORKBOOK_PATH.read_bytes()

    plan = prepare_generated_division_roster_export(
        division_layout=build.division,
        dataset=result.dataset,
        version=result.version,
        generated=result.generated,
    )

    assert canonical_json(result.version.model_dump(mode="json")) == version_before
    assert DIVISION_WORKBOOK_PATH.read_bytes() == template_before
    assert plan.report.publication_state == "blocked"
    assert plan.report.review_export_allowed is True
    assert not plan.report.reconciliation.errors
    assert not plan.report.export_failures


def test_prepared_plan_rejects_same_version_id_with_different_content(
    prepared_export_plan,
):
    build, result, plan = prepared_export_plan
    requested = plan.review_version.model_copy(deep=True)
    requested.entries[0].worker_name = "同 ID 但內容被竄改"

    with pytest.raises(ValueError, match="content does not match requested version"):
        build_generated_division_roster_workbook(
            template_path=DIVISION_WORKBOOK_PATH,
            division_layout=build.division,
            dataset=result.dataset,
            version=requested,
            prepared_plan=plan,
        )


def test_prepared_plan_recomputes_and_rejects_forged_content_hash(
    prepared_export_plan,
):
    build, result, original = prepared_export_plan
    plan = deepcopy(original)
    assert plan.review_version.reconciliation is not None
    plan.review_version.reconciliation.content_hash = "0" * 64
    plan.report.reconciliation.content_hash = "0" * 64

    with pytest.raises(ValueError, match="reconciliation content hash is invalid"):
        build_generated_division_roster_workbook(
            template_path=DIVISION_WORKBOOK_PATH,
            division_layout=build.division,
            dataset=result.dataset,
            version=plan.review_version,
            prepared_plan=plan,
        )


def test_prepared_plan_rejects_missing_needs_review_manifest_audits(
    prepared_export_plan,
):
    build, result, original = prepared_export_plan
    plan = deepcopy(original)
    manifest = next(
        item for item in plan.report.placements
        if item.disposition == "needs_review"
    )
    manifest.audit_ids = []

    with pytest.raises(ValueError, match="placement payload drifted"):
        build_generated_division_roster_workbook(
            template_path=DIVISION_WORKBOOK_PATH,
            division_layout=build.division,
            dataset=result.dataset,
            version=plan.review_version,
            prepared_plan=plan,
        )


def test_prepared_plan_rejects_target_entry_with_same_id_different_payload(
    prepared_export_plan,
):
    build, result, original = prepared_export_plan
    plan = deepcopy(original)
    target = plan.placements[0]
    target.entry = target.entry.model_copy(update={
        "worker_name": "同 entry ID 但 payload 被竄改",
    })

    with pytest.raises(ValueError, match="target entry payload drifted"):
        build_generated_division_roster_workbook(
            template_path=DIVISION_WORKBOOK_PATH,
            division_layout=build.division,
            dataset=result.dataset,
            version=plan.review_version,
            prepared_plan=plan,
        )


def test_prepared_plan_rejects_changed_cell_note_drift(prepared_export_plan):
    build, result, original = prepared_export_plan
    plan = deepcopy(original)
    ref = next(iter(plan.changed_cells))
    plan.changed_cells[ref].note = "竄改後的非 canonical comment"

    with pytest.raises(ValueError, match="changed-cell marker drifted"):
        build_generated_division_roster_workbook(
            template_path=DIVISION_WORKBOOK_PATH,
            division_layout=build.division,
            dataset=result.dataset,
            version=plan.review_version,
            prepared_plan=plan,
        )


def test_prepared_plan_integrity_hash_catches_unmodeled_report_drift(
    prepared_export_plan,
):
    build, result, original = prepared_export_plan
    plan = deepcopy(original)
    plan.report.export_block_reasons.append("竄改後的額外阻塞理由")

    with pytest.raises(ValueError, match="plan integrity hash is invalid"):
        build_generated_division_roster_workbook(
            template_path=DIVISION_WORKBOOK_PATH,
            division_layout=build.division,
            dataset=result.dataset,
            version=plan.review_version,
            prepared_plan=plan,
        )


def _embedded_reconciliation(ws) -> str:
    manifest_col = next(
        cell.column for cell in ws[1]
        if cell.value == "RC_reconciliation_key"
    )
    parts: list[tuple[str, str]] = []
    for row in range(2, ws.max_row + 1):
        key = ws.cell(row, manifest_col).value
        value = ws.cell(row, manifest_col + 1).value
        if isinstance(key, str) and key.startswith("reconciliation_json_"):
            parts.append((key, value or ""))
    return "".join(value for _, value in sorted(parts))


def _sheet_key_values(ws) -> dict[str, object]:
    return {
        str(ws.cell(row, 1).value): ws.cell(row, 2).value
        for row in range(2, ws.max_row + 1)
        if ws.cell(row, 1).value is not None
    }
